"""Parse CSV and Excel files into member import rows. Headers: company_id, scheme_id, card_no, name, dob, status."""

import csv
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.application.dtos.member import MemberCreate



def _normalize_header(h: str) -> str:
    return (h or "").strip().lower().replace(" ", "_").replace("-", "_")


def _parse_dob(value: Any) -> date | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


# First-class member fields (our names). Legacy CSV may use aliases below.
FIRST_CLASS_KEYS = frozenset({
    "company_id", "scheme_id", "card_no", "name", "dob", "status",
    "employee_no", "gender", "address", "tel_home", "tel_mobile", "email",
    "department", "branch", "occupation", "date_of_joining", "date_of_leaving", "remarks",
})
# Normalized legacy headers (Members.csv) -> our first-class key (so we don't put them in extra).
LEGACY_MEMBER_ALIASES = {
    "employeeno": "employee_no",
    "telhome": "tel_home",
    "telmobile": "tel_mobile",
    "dateofjoining": "date_of_joining",
    "dateofleaving": "date_of_leaving",
}


def _get_first_class(row: dict[str, Any], our_key: str, legacy_keys: list[str] | None = None) -> str:
    """Get string value from row by our key or legacy alias. Returns stripped or ''."""
    v = row.get(our_key)
    if v is None and legacy_keys:
        for k in legacy_keys:
            v = row.get(k)
            if v is not None:
                break
    return str(v).strip() if v is not None else ""


def _row_to_member_create(row: dict[str, Any]) -> tuple[MemberCreate | None, str | None]:
    """Convert a dict of column->value to MemberCreate. Returns (data, error_message)."""
    def get(key: str) -> str:
        v = row.get(key) or row.get(key.replace("_", " ")) or ""
        return str(v).strip() if v is not None else ""

    company_id = get("company_id")
    scheme_id = get("scheme_id")
    card_no = get("card_no")
    name = get("name")
    if not name:
        return None, "name is required"
    if not company_id:
        return None, "company_id is required"
    if not scheme_id:
        return None, "scheme_id is required"
    if not card_no:
        return None, "card_no is required"

    dob = _parse_dob(row.get("dob"))
    status = (get("status") or "active").strip() or "active"

    employee_no = _get_first_class(row, "employee_no", ["employeeno"]) or None
    gender = _get_first_class(row, "gender") or None
    address = _get_first_class(row, "address") or None
    tel_home = _get_first_class(row, "tel_home", ["telhome"]) or None
    tel_mobile = _get_first_class(row, "tel_mobile", ["telmobile"]) or None
    email = _get_first_class(row, "email") or None
    department = _get_first_class(row, "department") or None
    branch = _get_first_class(row, "branch") or None
    occupation = _get_first_class(row, "occupation") or None
    date_of_joining = _parse_dob(row.get("date_of_joining") or row.get("dateofjoining"))
    date_of_leaving = _parse_dob(row.get("date_of_leaving") or row.get("dateofleaving"))
    remarks = _get_first_class(row, "remarks") or None

    keys_used = set(FIRST_CLASS_KEYS) | set(LEGACY_MEMBER_ALIASES)
    extra: dict[str, Any] = {}
    for k, v in row.items():
        if k not in keys_used and v is not None and str(v).strip() != "":
            extra[k] = v
    if not extra:
        extra = None

    return (
        MemberCreate(
            company_id=company_id,
            scheme_id=scheme_id,
            card_no=card_no,
            name=name,
            dob=dob,
            status=status[:32],
            employee_no=employee_no[:64] if employee_no else None,
            gender=gender[:32] if gender else None,
            address=address[:255] if address else None,
            tel_home=tel_home[:64] if tel_home else None,
            tel_mobile=tel_mobile[:64] if tel_mobile else None,
            email=email[:255] if email else None,
            department=department[:255] if department else None,
            branch=branch[:255] if branch else None,
            occupation=occupation[:255] if occupation else None,
            date_of_joining=date_of_joining,
            date_of_leaving=date_of_leaving,
            remarks=remarks[:512] if remarks else None,
            extra=extra,
        ),
        None,
    )


def parse_csv(content: bytes) -> tuple[list[MemberCreate], list[int], list[tuple[int, str]]]:
    """
    Parse CSV bytes (first row = headers). Column names normalized (lowercase, spaces->underscores).
    Returns (items, file_row_for_item, parse_errors).
    - items: valid MemberCreate list
    - file_row_for_item: 1-based file row index for each item (so errors can reference file row)
    - parse_errors: (row_1based, message) for rows that failed parsing
    """
    reader = csv.DictReader(BytesIO(content).read().decode("utf-8-sig").splitlines())
    fieldnames = reader.fieldnames or []
    normalized = {_normalize_header(h) for h in fieldnames}
    required = {"company_id", "scheme_id", "card_no", "name"}
    missing = required - normalized
    if missing:
        return [], [], [(1, f"Missing required columns: {', '.join(sorted(missing))}")]

    items: list[MemberCreate] = []
    file_row_for_item: list[int] = []
    errors: list[tuple[int, str]] = []
    for i, raw_row in enumerate(reader):
        row = {}
        for orig, val in raw_row.items():
            norm = _normalize_header(orig)
            if norm:
                row[norm] = val
        create, err = _row_to_member_create(row)
        if err:
            errors.append((i + 2, err))  # 1-based + header row
        elif create:
            items.append(create)
            file_row_for_item.append(i + 2)
    return items, file_row_for_item, errors


def parse_excel(content: bytes) -> tuple[list[MemberCreate], list[int], list[tuple[int, str]]]:
    """
    Parse first sheet of xlsx. First row = headers. Returns (items, file_row_for_item, parse_errors).
    """
    wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    sheet = wb.active
    if not sheet:
        wb.close()
        return [], [], [(1, "Workbook has no active sheet")]
    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        wb.close()
        return [], [], [(1, "Sheet is empty")]
    headers = [_normalize_header(str(h)) for h in header_row]
    col_index: dict[str, int] = {h: i for i, h in enumerate(headers) if h}
    required = {"company_id", "scheme_id", "card_no", "name"}
    missing = required - set(col_index)
    if missing:
        wb.close()
        return [], [], [(1, f"Missing required columns: {', '.join(sorted(missing))}")]

    items: list[MemberCreate] = []
    file_row_for_item: list[int] = []
    errors: list[tuple[int, str]] = []
    for row_index, row_tuple in enumerate(rows_iter):
        row = {h: (row_tuple[col_index[h]] if col_index[h] < len(row_tuple) else None) for h in col_index}
        create, err = _row_to_member_create(row)
        if err:
            errors.append((row_index + 2, err))  # 1-based + header row
        elif create:
            items.append(create)
            file_row_for_item.append(row_index + 2)
    wb.close()
    return items, file_row_for_item, errors
